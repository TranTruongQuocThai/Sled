import json
import unicodedata
from collections import defaultdict
from io import BytesIO
from zipfile import BadZipFile, ZipFile
from xml.etree.ElementTree import ParseError

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from models.repository import get_db, record_audit
from models.workflows import atomic, insert_rows, now, quantity


def normalize(value):
    value=str(value or '').strip().lower().replace('đ','d')
    return ''.join(c for c in unicodedata.normalize('NFD',value) if unicodedata.category(c)!='Mn')


def value(row,names):
    for name in names:
        cell=row.get(name)
        if cell is not None and str(cell).strip():
            return str(cell).strip()
    return ''


def read_rows(file_storage):
    workbook=None
    try:
        data=file_storage.read(8*1024*1024+1)
        if len(data)>8*1024*1024:
            raise ValueError('File Excel tối đa 8 MB.')
        with ZipFile(BytesIO(data)) as archive:
            if sum(i.file_size for i in archive.infolist())>64*1024*1024:
                raise ValueError('File Excel sau giải nén quá lớn.')
        workbook=load_workbook(BytesIO(data),read_only=True,data_only=True)
        sheet=workbook.active
        if sheet is None or not hasattr(sheet,'iter_rows'):
            raise ValueError('Workbook không có trang dữ liệu hợp lệ.')
        if (sheet.max_row or 0)>10001 or (sheet.max_column or 0)>32:
            raise ValueError('Mỗi file tối đa 10.000 dòng dữ liệu và 32 cột.')
        iterator=sheet.iter_rows(values_only=True)
        headers=[normalize(c) for c in next(iterator,())]
        present=[h for h in headers if h]
        if not present or len(set(present))!=len(present):
            raise ValueError('Tiêu đề cột bị thiếu hoặc trùng.')
        rows=[]
        for row_number,cells in enumerate(iterator,2):
            if len(cells)>32:
                raise ValueError('Mỗi file tối đa 32 cột.')
            if any(c is not None and str(c).strip() for c in cells):
                rows.append((row_number,dict(zip(headers,cells))))
            if row_number>10001:
                raise ValueError('Mỗi file tối đa 10.000 dòng.')
        if not rows:
            raise ValueError('File Excel không có dòng dữ liệu.')
        return rows
    except (BadZipFile, InvalidFileException, ParseError, OSError, KeyError, TypeError, RuntimeError) as error:
        raise ValueError('Không đọc được file. Hãy chọn file Excel .xlsx hợp lệ.') from error
    finally:
        if workbook:
            workbook.close()


def preview_import(file_storage,import_type,user_id):
    if import_type not in ('components','bom'):
        raise ValueError('Loại import không hợp lệ.')
    rows=read_rows(file_storage)
    payloads=[]
    seen=set()
    db=get_db()
    components={r['sku'].upper():dict(r) for r in db.execute('SELECT id,sku,unit,active FROM components')}
    products={r['sku'].upper():dict(r) for r in db.execute('SELECT id,sku,active FROM products')}
    for row_number,row in rows:
        error=None
        if import_type=='components':
            payload={'sku':value(row,('sku','ma linh kien')).upper(),
                     'name':value(row,('name','ten linh kien')),
                     'unit':value(row,('unit','don vi tinh','don vi','dvt')),
                     'category':value(row,('category','nhom linh kien','nhom'))}
            if not all(payload[k] for k in ('sku','name','unit')):
                error='Thiếu mã, tên hoặc đơn vị tính linh kiện.'
            duplicate=payload['sku']
            existing=components.get(payload['sku'])
            payload['existing']=bool(existing)
            if existing and existing['unit']!=payload['unit']:
                error='Đơn vị tính khác danh mục hiện có; không tự quy đổi khi import.'
        else:
            payload={'product_sku':value(row,('product_sku','product sku','ma san pham')).upper(),
                     'component_sku':value(row,('component_sku','component sku','ma linh kien')).upper(),
                     'quantity':value(row,('quantity','dinh muc','so luong'))}
            try:
                payload['quantity']=quantity(payload['quantity'])
            except ValueError as exc:
                error=str(exc)
            if not payload['product_sku'] or not payload['component_sku']:
                error='Thiếu mã sản phẩm hoặc linh kiện.'
            if not products.get(payload['product_sku'],{}).get('active'):
                error='Mã sản phẩm không tồn tại hoặc đã ngừng sử dụng.'
            if not components.get(payload['component_sku'],{}).get('active'):
                error='Mã linh kiện không tồn tại hoặc đã ngừng sử dụng.'
            duplicate=(payload['product_sku'],payload['component_sku'])
        if duplicate in seen:
            error='Trùng mã trong file.'
        seen.add(duplicate)
        payloads.append((row_number,payload,error))
    with atomic() as db:
        db.execute('INSERT INTO import_batches(import_type,created_by,created_at) VALUES(?,?,?)',(import_type,user_id,now()))
        batch_id=db.execute('SELECT last_insert_rowid()').fetchone()[0]
        insert_rows(db,'INSERT INTO import_batch_rows(batch_id,row_number,payload,error)',
            [(batch_id,n,json.dumps(payload,ensure_ascii=False),error) for n,payload,error in payloads])
    return batch_id,payloads,sum(bool(error) for _,_,error in payloads)


def commit_import(batch_id,import_type,user_id,mode='create'):
    if import_type not in ('components','bom') or mode not in ('create','upsert'):
        raise ValueError('Loại import hoặc chế độ cập nhật không hợp lệ.')
    with atomic() as db:
        batch=db.execute("SELECT * FROM import_batches WHERE id=? AND import_type=? AND created_by=? AND status='preview'",(batch_id,import_type,user_id)).fetchone()
        if not batch:
            raise ValueError('Phiên import không tồn tại hoặc đã xác nhận.')
        rows=db.execute('SELECT * FROM import_batch_rows WHERE batch_id=? ORDER BY row_number',(batch_id,)).fetchall()
        if not rows or any(row['error'] for row in rows):
            raise ValueError('File rỗng hoặc còn dòng lỗi; chưa nhập dữ liệu.')
        components={r['sku'].upper():dict(r) for r in db.execute('SELECT id,sku,unit,active FROM components')}
        products={r['sku'].upper():dict(r) for r in db.execute('SELECT id,sku,active FROM products')}
        payloads=[json.loads(row['payload']) for row in rows]
        if import_type=='components':
            for payload in payloads:
                existing=components.get(payload['sku'])
                if existing and mode=='create':
                    raise ValueError(f"Mã {payload['sku']} đã tồn tại. Chọn chế độ cập nhật hoặc sửa file.")
                if existing and existing['unit']!=payload['unit']:
                    raise ValueError(f"Đơn vị {payload['sku']} đã thay đổi sau khi xem trước.")
            insert_rows(db,'INSERT OR IGNORE INTO component_categories(name)',
                        [(name,) for name in sorted({p['category'] for p in payloads if p['category']})])
            categories={r['name']:r['id'] for r in db.execute('SELECT id,name FROM component_categories')}
            insert_rows(db,'INSERT INTO components(sku,name,unit,category_id)',
                [(components[p['sku']]['sku'] if p['sku'] in components else p['sku'],p['name'],p['unit'],categories.get(p['category'])) for p in payloads],
                ' ON CONFLICT(sku) DO UPDATE SET name=excluded.name,category_id=excluded.category_id' if mode=='upsert' else '')
        else:
            groups=defaultdict(list)
            for row,payload in zip(rows,payloads):
                product=products.get(payload['product_sku'])
                component=components.get(payload['component_sku'])
                if not product or not product['active'] or not component or not component['active']:
                    raise ValueError(f"Sản phẩm hoặc linh kiện không còn hợp lệ tại dòng {row['row_number']}.")
                groups[product['id']].append((component['id'],quantity(payload['quantity'])))
            versions={r['product_id']:r['version']+1 for r in db.execute('SELECT product_id,MAX(version) version FROM boms GROUP BY product_id')}
            insert_rows(db,'INSERT INTO boms(product_id,version,status,created_by,created_at,import_batch_id)',
                        [(pid,versions.get(pid,1),'draft',user_id,now(),batch_id) for pid in groups])
            ids={r['product_id']:r['id'] for r in db.execute("SELECT id,product_id,version FROM boms WHERE status='draft'")
                 if r['product_id'] in groups and r['version']==versions.get(r['product_id'],1)}
            insert_rows(db,'INSERT INTO bom_items(bom_id,component_id,quantity)',
                        [(ids[pid],cid,count) for pid,items in groups.items() for cid,count in items])
        db.execute("UPDATE import_batches SET status='committed' WHERE id=?",(batch_id,))
        record_audit(user_id,'excel_import_committed',import_type,batch_id)
    return list(ids.values()) if import_type=='bom' else []
